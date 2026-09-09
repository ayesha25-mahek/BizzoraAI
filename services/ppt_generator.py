import json
import os

from pptx import Presentation
from pptx.util import Inches, Pt


class PPTGenerator:

    def __init__(self):

        self.output_dir = "output/presentations"

        os.makedirs(
            self.output_dir,
            exist_ok=True
        )

    def generate(
        self,
        slide_plan,
        filename="bizzoraai_presentation.pptx"
    ):

        presentation = Presentation()

        # Widescreen format
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)

        for slide_data in slide_plan:

            slide_type = slide_data.get(
                "type",
                "content"
            )

            title = slide_data.get(
                "title",
                ""
            )

            content = slide_data.get(
                "content",
                []
            )

            if slide_type == "title":

                self._create_title_slide(
                    presentation,
                    title,
                    content
                )

            elif slide_type == "overview":

                self._create_content_slide(
                    presentation,
                    title,
                    content
                )

            elif slide_type == "key_points":

                self._create_bullet_slide(
                    presentation,
                    title,
                    content
                )

            elif slide_type == "statistics":

                self._create_statistics_slide(
                    presentation,
                    title,
                    content
                )

            elif slide_type == "recommendations":

                self._create_bullet_slide(
                    presentation,
                    title,
                    content
                )

            elif slide_type == "conclusion":

                self._create_conclusion_slide(
                    presentation,
                    title,
                    content
                )

            else:

                self._create_content_slide(
                    presentation,
                    title,
                    content
                )

        output_path = os.path.join(
            self.output_dir,
            filename
        )

        presentation.save(
            output_path
        )

        print(
            f"\n[SUCCESS] PowerPoint generated:"
            f"\n{output_path}"
        )

        return output_path

    # --------------------------------------------------
    # TITLE SLIDE
    # --------------------------------------------------

    def _create_title_slide(
        self,
        presentation,
        title,
        content
    ):

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[6]
        )

        title_box = slide.shapes.add_textbox(
            Inches(1),
            Inches(2.2),
            Inches(11.3),
            Inches(1.5)
        )

        title_frame = title_box.text_frame

        title_frame.text = title

        paragraph = title_frame.paragraphs[0]

        paragraph.font.size = Pt(36)
        paragraph.font.bold = True

        subtitle = ""

        if isinstance(content, str):

            subtitle = content

        elif isinstance(content, list) and content:

            subtitle = str(content[0])

        subtitle_box = slide.shapes.add_textbox(
            Inches(1),
            Inches(4),
            Inches(11.3),
            Inches(1.5)
        )

        subtitle_frame = subtitle_box.text_frame

        subtitle_frame.text = subtitle

        subtitle_frame.paragraphs[0].font.size = Pt(20)

    # --------------------------------------------------
    # CONTENT SLIDE
    # --------------------------------------------------

    def _create_content_slide(
        self,
        presentation,
        title,
        content
    ):

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[6]
        )

        self._add_title(
            slide,
            title
        )

        text = self._normalise_content(
            content
        )

        box = slide.shapes.add_textbox(
            Inches(1),
            Inches(1.8),
            Inches(11.3),
            Inches(4.8)
        )

        frame = box.text_frame

        frame.text = text

        frame.paragraphs[0].font.size = Pt(22)

    # --------------------------------------------------
    # BULLET SLIDE
    # --------------------------------------------------

    def _create_bullet_slide(
        self,
        presentation,
        title,
        content
    ):

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[6]
        )

        self._add_title(
            slide,
            title
        )

        box = slide.shapes.add_textbox(
            Inches(1),
            Inches(1.8),
            Inches(11.3),
            Inches(4.8)
        )

        frame = box.text_frame

        frame.clear()

        if isinstance(content, str):

            content = [content]

        for index, item in enumerate(content):

            paragraph = (
                frame.paragraphs[0]
                if index == 0
                else frame.add_paragraph()
            )

            paragraph.text = str(item)

            paragraph.level = 0

            paragraph.font.size = Pt(22)

    # --------------------------------------------------
    # STATISTICS
    # --------------------------------------------------

    def _create_statistics_slide(
        self,
        presentation,
        title,
        content
    ):

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[6]
        )

        self._add_title(
            slide,
            title
        )

        if isinstance(content, str):

            content = [content]

        count = len(content)

        if count == 0:

            return

        width = 10.5 / min(count, 3)

        for index, statistic in enumerate(
            content[:3]
        ):

            x = 1 + (
                index * width
            )

            box = slide.shapes.add_textbox(
                Inches(x),
                Inches(2.5),
                Inches(width - 0.3),
                Inches(2)
            )

            frame = box.text_frame

            frame.text = str(statistic)

            paragraph = frame.paragraphs[0]

            paragraph.font.size = Pt(24)
            paragraph.font.bold = True

    # --------------------------------------------------
    # CONCLUSION
    # --------------------------------------------------

    def _create_conclusion_slide(
        self,
        presentation,
        title,
        content
    ):

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[6]
        )

        self._add_title(
            slide,
            title
        )

        text = self._normalise_content(
            content
        )

        box = slide.shapes.add_textbox(
            Inches(1),
            Inches(2.3),
            Inches(11.3),
            Inches(3)
        )

        frame = box.text_frame

        frame.text = text

        frame.paragraphs[0].font.size = Pt(26)

    # --------------------------------------------------
    # TITLE
    # --------------------------------------------------

    def _add_title(
        self,
        slide,
        title
    ):

        box = slide.shapes.add_textbox(
            Inches(0.8),
            Inches(0.5),
            Inches(11.7),
            Inches(1)
        )

        frame = box.text_frame

        frame.text = title

        paragraph = frame.paragraphs[0]

        paragraph.font.size = Pt(30)
        paragraph.font.bold = True

    # --------------------------------------------------
    # NORMALISE CONTENT
    # --------------------------------------------------

    def _normalise_content(
        self,
        content
    ):

        if isinstance(content, str):

            return content

        if isinstance(content, list):

            return "\n\n".join(
                str(item)
                for item in content
            )

        return str(content)