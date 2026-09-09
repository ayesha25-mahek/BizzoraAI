import os
from pathlib import Path

from pypdf import PdfReader
from docx import Document
from pptx import Presentation


class DocumentReader:

    def read(self, file_path):

        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(
                f"File not found: {file_path}"
            )

        extension = path.suffix.lower()

        print(f"\n📄 Reading {extension} file...")

        if extension == ".pdf":

            return self._read_pdf(path)

        elif extension == ".docx":

            return self._read_docx(path)

        elif extension == ".txt":

            return self._read_txt(path)

        elif extension == ".pptx":

            return self._read_pptx(path)

        else:

            raise ValueError(
                f"Unsupported file type: {extension}"
            )


    # --------------------------------------------------
    # PDF
    # --------------------------------------------------

    def _read_pdf(self, path):

        reader = PdfReader(str(path))

        pages = []

        for page_number, page in enumerate(
            reader.pages,
            start=1
        ):

            text = page.extract_text()

            if text:

                pages.append(
                    f"\n--- Page {page_number} ---\n"
                    + text
                )

        content = "\n".join(pages)

        print(
            f"[SUCCESS] PDF read successfully "
            f"({len(reader.pages)} pages)"
        )

        return content


    # --------------------------------------------------
    # DOCX
    # --------------------------------------------------

    def _read_docx(self, path):

        document = Document(str(path))

        paragraphs = []

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:

                paragraphs.append(text)

        content = "\n".join(paragraphs)

        print("[SUCCESS] DOCX read successfully")

        return content


    # --------------------------------------------------
    # TXT
    # --------------------------------------------------

    def _read_txt(self, path):

        with open(
            path,
            "r",
            encoding="utf-8"
        ) as file:

            content = file.read()

        print("[SUCCESS] TXT read successfully")

        return content


    # --------------------------------------------------
    # PPTX
    # --------------------------------------------------

    def _read_pptx(self, path):

        presentation = Presentation(str(path))

        slides = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1
        ):

            slide_text = []

            for shape in slide.shapes:

                if hasattr(shape, "text"):

                    text = shape.text.strip()

                    if text:

                        slide_text.append(text)

            if slide_text:

                slides.append(
                    f"\n--- Slide {slide_number} ---\n"
                    + "\n".join(slide_text)
                )

        content = "\n".join(slides)

        print(
            f"[SUCCESS] PPTX read successfully "
            f"({len(presentation.slides)} slides)"
        )

        return content